from __future__ import annotations

import json
import re
import shutil
from dataclasses import dataclass, field
from datetime import datetime
from html import escape, unescape
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import unquote, urlparse

from app.schemas import TranscriptSegment
from app.services.httpclient import http_client
from app.services.ingest.base import DOCUMENT_EXTS, path_suffix
from app.services.sourcetime import parse_source_datetime, pick_html_datetime
from app.services.textnorm import normalize_transcript

DEFAULT_UA = (
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
    "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/128.0.0.0 Safari/537.36"
)
_JS_STRING_RE = re.compile(r"\\u([0-9a-fA-F]{4})|\\(.)")
_H5_ARTICLE_RE = re.compile(r"^/h5/article/([^/]+)/?$")
_ARTICLE_EMBED_RE = re.compile(
    r'\barticle_id\b[\s\S]{0,800}?title:"((?:\\.|[^"\\])+)"[\s\S]{0,4000}?content:"((?:\\.|[^"\\])+)"',
)
_JSONLD_RE = re.compile(
    r'<script[^>]*type=["\']application/ld\+json["\'][^>]*>(.*?)</script>',
    re.I | re.S,
)


class DocumentError(RuntimeError):
    """文档提取失败。"""


SCAN_CHARS_PER_PAGE = 40
MAX_DOWNLOAD_BYTES = 80 * 1024 * 1024
CHUNK_CHARS = 700

CONTENT_EXT = {
    "application/pdf": ".pdf",
    "application/msword": ".doc",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/html": ".html",
    "text/plain": ".txt",
    "text/markdown": ".md",
    "application/xhtml+xml": ".html",
}


@dataclass
class ExtractedDocument:
    segments: list[TranscriptSegment]
    title: str = ""
    author: str = ""
    suffix: str = ""
    used_ocr: bool = False
    created_at: datetime | None = None
    extra: dict = field(default_factory=dict)


def looks_scanned(pages: list[str]) -> bool:
    nonempty = [item.strip() for item in pages if item and item.strip()]
    if not pages:
        return True
    if not nonempty:
        return True
    total = sum(len(item) for item in nonempty)
    return total / max(len(pages), 1) < SCAN_CHARS_PER_PAGE


def extract_path(path: Path, *, ocr=None, convert_doc=None, progress=None) -> ExtractedDocument:
    if not path.exists() or not path.is_file():
        raise DocumentError(f"找不到文档文件：{path}")
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return extract_pdf(path, ocr=ocr, progress=progress)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".doc":
        converter = convert_doc or convert_legacy_doc
        if converter is convert_legacy_doc:
            from app.services.plugins import convert_doc_with_soffice, ensure_plugin

            ensure_plugin("legacy_doc", progress=progress)
            converted = convert_doc_with_soffice(path, path.parent)
        else:
            converted = converter(path)
        return extract_docx(converted)
    if suffix in {".md", ".markdown", ".txt"}:
        return extract_plain(path)
    if suffix in {".html", ".htm"}:
        return extract_html_file(path)
    raise DocumentError(f"暂不支持该文档格式：{suffix or path.name}")


def candidate_page_urls(url: str) -> list[str]:
    raw = (url or "").strip()
    if not raw:
        return []
    found = [raw]
    parsed = urlparse(raw)
    path = parsed.path or ""
    match = _H5_ARTICLE_RE.match(path)
    if match:
        found.append(parsed._replace(path=f"/a/{match.group(1)}").geturl())
    if path.startswith("/h5/") and len(path) > 4:
        found.append(parsed._replace(path=path[3:]).geturl())
    unique: list[str] = []
    for item in found:
        if item not in unique:
            unique.append(item)
    return unique


def extract_url(url: str, dest_dir: Path, headers: dict[str, str] | None = None, progress=None) -> ExtractedDocument:
    dest_dir.mkdir(parents=True, exist_ok=True)
    last_error: Exception | None = None
    for candidate in candidate_page_urls(url):
        try:
            path, header_date = download_url(candidate, dest_dir, headers=headers or {})
            raw = path.read_bytes()
            if _looks_html(raw) and path.suffix.lower() not in {".pdf", ".docx", ".doc"}:
                extracted = extract_html_bytes(
                    raw, fallback_title=Path(urlparse(candidate).path).stem or path.stem
                )
            else:
                extracted = extract_path(path, progress=progress)
            if extracted.created_at is None:
                extracted.created_at = header_date
            extracted.extra["source_path"] = str(path)
            return extracted
        except DocumentError as exc:
            last_error = exc
    raise last_error or DocumentError("未能从网页提取到正文")


def download_url(url: str, dest_dir: Path, headers: dict[str, str] | None = None) -> tuple[Path, datetime | None]:
    dest_dir.mkdir(parents=True, exist_ok=True)
    request_headers = {"User-Agent": DEFAULT_UA, **(headers or {})}
    with http_client(timeout=60.0, headers=request_headers, follow_redirects=True) as client:
        try:
            response = client.get(url)
            response.raise_for_status()
        except Exception as exc:
            raise DocumentError(f"下载文档失败：{exc}") from exc
    data = response.content
    if len(data) > MAX_DOWNLOAD_BYTES:
        raise DocumentError("文档超过 80MB，请缩小后再导入")
    suffix = path_suffix(url)
    content_type = (response.headers.get("content-type") or "").split(";")[0].strip().lower()
    if suffix not in DOCUMENT_EXTS:
        suffix = CONTENT_EXT.get(content_type, "")
        if not suffix:
            suffix = _sniff_suffix(data)
    if not suffix:
        suffix = ".bin"
    dest = dest_dir / f"source{suffix}"
    dest.write_bytes(data)
    header_date = parse_source_datetime(response.headers.get("last-modified"))
    return dest, header_date


def extract_pdf(path: Path, *, ocr=None, progress=None) -> ExtractedDocument:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentError("缺少 pymupdf，无法读取 PDF") from exc
    try:
        doc = fitz.open(path)
    except Exception as exc:
        raise DocumentError(f"无法打开 PDF：{exc}") from exc
    pages: list[str] = []
    meta = {}
    try:
        meta = doc.metadata or {}
        for page in doc:
            pages.append(normalize_transcript(page.get_text("text") or ""))
    finally:
        doc.close()
    used_ocr = False
    if looks_scanned(pages):
        if ocr is None:
            from app.services.plugins import ensure_ocr_engine

            ocr = ensure_ocr_engine(progress=progress)
        pages = ocr_pdf_pages(path, ocr)
        used_ocr = True
    segments = _page_segments(pages)
    if not segments:
        raise DocumentError("未能提取到文字（可能是空白页或扫描件识别失败）")
    created = parse_source_datetime(meta.get("creationDate")) or parse_source_datetime(meta.get("modDate"))
    return ExtractedDocument(
        segments=segments,
        title=path.stem,
        suffix=".pdf",
        used_ocr=used_ocr,
        created_at=created,
    )


def ocr_pdf_pages(path: Path, ocr) -> list[str]:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentError("缺少 pymupdf，无法渲染扫描 PDF") from exc
    doc = fitz.open(path)
    pages: list[str] = []
    try:
        for index, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            image_path = path.parent / f".ocr-{path.stem}-{index}.png"
            try:
                pix.save(str(image_path))
                pages.append(run_ocr(ocr, image_path))
            finally:
                image_path.unlink(missing_ok=True)
    finally:
        doc.close()
    return pages


def run_ocr(ocr, image_path: Path) -> str:
    try:
        result = ocr(str(image_path))
    except Exception as exc:
        raise DocumentError(f"OCR 识别失败：{exc}") from exc
    payload = result[0] if isinstance(result, tuple) else result
    lines: list[str] = []
    if not payload:
        return ""
    for item in payload:
        if isinstance(item, str):
            lines.append(item)
            continue
        if isinstance(item, (list, tuple)) and len(item) >= 2:
            text = item[1]
            if isinstance(text, str):
                lines.append(text)
            elif isinstance(text, (list, tuple)) and text:
                lines.append(str(text[0]))
    return normalize_transcript("".join(lines))


def extract_docx(path: Path) -> ExtractedDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentError("缺少 python-docx，无法读取 Word 文档") from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise DocumentError(f"无法打开 Word 文档：{exc}") from exc
    paragraphs = [normalize_transcript(item.text or "") for item in document.paragraphs]
    paragraphs = [item for item in paragraphs if item]
    segments = chunk_paragraphs(paragraphs)
    if not segments:
        raise DocumentError("Word 文档里没有可提取的文字")
    core = document.core_properties
    created = parse_source_datetime(core.created) or parse_source_datetime(core.modified)
    return ExtractedDocument(
        segments=segments,
        title=(core.title or "").strip() or path.stem,
        author=(core.author or "").strip(),
        suffix=path.suffix.lower(),
        created_at=created,
    )


def convert_legacy_doc(path: Path, dest_dir: Path | None = None) -> Path:
    from app.services.plugins import convert_doc_with_soffice, ensure_plugin

    ensure_plugin("legacy_doc")
    return convert_doc_with_soffice(path, dest_dir or path.parent)


def extract_plain(path: Path) -> ExtractedDocument:
    text = path.read_text(encoding="utf-8", errors="ignore")
    segments = chunk_paragraphs(_split_blocks(text))
    if not segments:
        raise DocumentError("文本文件是空的")
    return ExtractedDocument(segments=segments, title=path.stem, suffix=path.suffix.lower())


def extract_html_file(path: Path) -> ExtractedDocument:
    return extract_html_bytes(path.read_bytes(), fallback_title=path.stem)


def document_file_created_at(path: Path) -> datetime | None:
    if not path.is_file():
        return None
    suffix = path.suffix.lower()
    if suffix in {".html", ".htm"}:
        return pick_html_datetime(path.read_text(encoding="utf-8", errors="ignore"))
    if suffix == ".docx":
        try:
            extracted = extract_docx(path)
        except Exception:
            return None
        return extracted.created_at
    if suffix == ".pdf":
        try:
            import fitz
        except ImportError:
            return None
        try:
            doc = fitz.open(path)
            try:
                meta = doc.metadata or {}
            finally:
                doc.close()
        except Exception:
            return None
        return parse_source_datetime(meta.get("creationDate")) or parse_source_datetime(meta.get("modDate"))
    return None


def extract_html_bytes(raw: bytes, fallback_title: str = "") -> ExtractedDocument:
    html = raw.decode("utf-8", errors="ignore")
    title, author, text = _embedded_article(html)
    created = pick_html_datetime(html)
    if not text:
        try:
            import trafilatura
        except ImportError as exc:
            raise DocumentError("缺少 trafilatura，无法提取网页正文") from exc
        text = trafilatura.extract(html, include_comments=False, include_tables=True) or ""
        metadata = trafilatura.extract_metadata(html)
        if metadata is not None:
            title = title or (metadata.title or "").strip()
            author = author or (metadata.author or "").strip()
            created = created or parse_source_datetime(getattr(metadata, "date", None))
    if not text:
        text = _html_to_text(html)
    title = title or fallback_title
    if text.lstrip().startswith("<"):
        text = _html_to_text(text)
    segments = chunk_paragraphs(_split_blocks(text))
    if not segments:
        raise DocumentError("未能从网页提取到正文")
    return ExtractedDocument(
        segments=segments, title=title, author=author, suffix=".html", created_at=created
    )


def _embedded_article(html: str) -> tuple[str, str, str]:
    title, author, text = _jsonld_article(html)
    if text:
        return title, author, text
    match = _ARTICLE_EMBED_RE.search(html)
    if not match:
        return "", "", ""
    title = _decode_js_string(match.group(1)).strip()
    content = _decode_js_string(match.group(2)).strip()
    return title, "", content


def _jsonld_article(html: str) -> tuple[str, str, str]:
    for block in _JSONLD_RE.findall(html):
        try:
            payload = json.loads(block)
        except json.JSONDecodeError:
            continue
        for node in _walk_jsonld(payload):
            types = node.get("@type") or node.get("type") or ""
            if isinstance(types, list):
                kinds = {str(item).lower() for item in types}
            else:
                kinds = {str(types).lower()}
            if not kinds.intersection({"article", "newsarticle", "blogposting"}):
                continue
            title = str(node.get("headline") or node.get("name") or "").strip()
            text = str(node.get("articleBody") or node.get("text") or "").strip()
            author = _ld_name(node.get("author"))
            if text:
                return title, author, text
    return "", "", ""


def _walk_jsonld(node):
    if isinstance(node, list):
        for item in node:
            yield from _walk_jsonld(item)
        return
    if not isinstance(node, dict):
        return
    yield node
    graph = node.get("@graph")
    if graph is not None:
        yield from _walk_jsonld(graph)
    for key, value in node.items():
        if key == "@graph":
            continue
        if isinstance(value, (dict, list)):
            yield from _walk_jsonld(value)


def _ld_name(value) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, dict):
        return str(value.get("name") or "").strip()
    if isinstance(value, list) and value:
        return _ld_name(value[0])
    return ""


def _decode_js_string(value: str) -> str:
    mapping = {"n": "\n", "r": "\r", "t": "\t", '"': '"', "'": "'", "\\": "\\", "/": "/"}

    def repl(match: re.Match[str]) -> str:
        if match.group(1):
            return chr(int(match.group(1), 16))
        return mapping.get(match.group(2), match.group(2) or "")

    return unescape(_JS_STRING_RE.sub(repl, value or ""))


class _HTMLText(HTMLParser):
    skip_tags = {"script", "style", "noscript", "svg", "head", "title", "meta", "link", "nav"}

    def __init__(self) -> None:
        super().__init__()
        self._skip = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        name = tag.lower()
        if name in self.skip_tags:
            self._skip += 1
        elif name in {"p", "div", "br", "li", "h1", "h2", "h3", "h4", "tr", "section", "blockquote"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        name = tag.lower()
        if name in self.skip_tags and self._skip:
            self._skip -= 1
        elif name in {"p", "div", "li", "h1", "h2", "h3", "h4", "tr", "section", "blockquote"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip:
            return
        text = unescape(data).strip()
        if text:
            self.parts.append(text)


def _html_to_text(html: str) -> str:
    parser = _HTMLText()
    try:
        parser.feed(html)
        parser.close()
    except Exception:
        return ""
    return normalize_transcript("".join(parser.parts))


def chunk_paragraphs(paragraphs: list[str], locator_kind: str = "段") -> list[TranscriptSegment]:
    blocks: list[str] = []
    buf: list[str] = []
    chars = 0
    for item in paragraphs:
        text = normalize_transcript(item)
        if not text:
            continue
        if chars and chars + len(text) > CHUNK_CHARS:
            blocks.append("\n".join(buf))
            buf = [text]
            chars = len(text)
            continue
        buf.append(text)
        chars += len(text)
    if buf:
        blocks.append("\n".join(buf))
    return [
        TranscriptSegment(id=index, start=0, end=0, text=block, locator=f"第{index + 1}{locator_kind}")
        for index, block in enumerate(blocks)
        if block.strip()
    ]


def _page_segments(pages: list[str]) -> list[TranscriptSegment]:
    segments: list[TranscriptSegment] = []
    for index, raw in enumerate(pages, start=1):
        text = normalize_transcript(raw)
        if not text:
            continue
        segments.append(TranscriptSegment(id=len(segments), start=0, end=0, text=text, locator=f"第{index}页"))
    return segments


def _split_blocks(text: str) -> list[str]:
    chunks = []
    normalized = (text or "").replace("\r\n", "\n").replace("\r", "\n")
    for block in normalized.split("\n"):
        cleaned = normalize_transcript(block)
        if cleaned:
            chunks.append(cleaned)
    if not chunks and (text or "").strip():
        chunks.append(normalize_transcript(text))
    return chunks


def _looks_html(data: bytes) -> bool:
    head = data[:200].lstrip().lower()
    return head.startswith(b"<!doctype") or head.startswith(b"<html") or b"<html" in head


def _sniff_suffix(data: bytes) -> str:
    if data.startswith(b"%PDF"):
        return ".pdf"
    if data.startswith(b"PK"):
        return ".docx"
    if data.startswith(b"\xd0\xcf\x11\xe0"):
        return ".doc"
    if _looks_html(data):
        return ".html"
    return ".txt"


def guessed_filename(url: str) -> str:
    name = Path(unquote(urlparse(url).path)).name
    return name or "source"


PREVIEW_TYPES = {
    ".pdf": "application/pdf",
    ".html": "text/html; charset=utf-8",
    ".htm": "text/html; charset=utf-8",
    ".txt": "text/plain; charset=utf-8",
    ".md": "text/plain; charset=utf-8",
    ".markdown": "text/plain; charset=utf-8",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
}


def ensure_job_source_file(job_id: str, original: Path) -> Path:
    from app.config import settings

    dest_dir = settings.job_workdir(job_id)
    suffix = original.suffix.lower() or ".bin"
    dest = dest_dir / f"source{suffix}"
    if dest.exists() and dest.stat().st_size > 0:
        try:
            if dest.resolve() == original.resolve():
                return dest
        except OSError:
            return dest
        return dest
    if not original.is_file() or original.stat().st_size <= 0:
        if dest.is_file() and dest.stat().st_size > 0:
            return dest
        raise DocumentError(f"找不到文档文件：{original}")
    try:
        if dest.resolve() == original.resolve():
            return dest
    except OSError:
        pass
    shutil.copy2(original, dest)
    return dest


def resolve_preview_file(job_id: str, source_path: str = "") -> Path:
    from app.config import settings

    workdir = settings.job_workdir(job_id).resolve()
    candidates: list[Path] = []
    if (source_path or "").strip():
        candidates.append(Path(source_path))
    candidates.extend(sorted(workdir.glob("source.*"), key=lambda item: item.stat().st_mtime, reverse=True))
    for path in candidates:
        if not path.exists() or not path.is_file() or path.stat().st_size <= 0:
            continue
        if path.name.startswith("."):
            continue
        try:
            resolved = path.resolve()
        except OSError:
            continue
        try:
            resolved.relative_to(workdir)
            return resolved
        except ValueError:
            return ensure_job_source_file(job_id, resolved)
    raise DocumentError("没有可预览的原文件")


_SCRIPT_RE = re.compile(r"<script\b[^>]*>[\s\S]*?</script>", re.I)
_STYLE_BLOCK_RE = re.compile(r"<style\b[^>]*>[\s\S]*?</style>", re.I)
_HTML_TAG_RE = re.compile(r"<[^>]+>")
_IMG_TAG_RE = re.compile(r"<img\b[^>]*>", re.I)
_ARTICLE_HINT_RE = re.compile(
    r"<(?:div|article|section|main)\b[^>]*(?:id|class)=[\"'][^\"']*"
    r"(?:js_content|rich_media_content|article-content|entry-content|"
    r"post-content|Post-RichText|article-body|post_content)[^\"']*[\"']",
    re.I,
)
_SPA_SHELL_RE = re.compile(
    r"""<div[^>]+id=["'](?:app|root|__nuxt|__next)["'][^>]*>\s*</div>""",
    re.I,
)
_ORIGINAL_PREVIEW_STYLE = """<meta name="referrer" content="no-referrer" />
<style id="vs-original-preview">
#js_content,.rich_media_content,.rich_media_area_primary,#img-content,#page-content{
  visibility:visible!important;opacity:1!important;display:block!important;height:auto!important;
}
html,body{overflow:auto!important;height:auto!important;}
#js_pc_qr_code,.qr_code_pc,#js_profile_card_modal,#js_alert_panel,#js_minipro_dialog,
#js_link_dialog,#js_product_dialog,#js_emotion_panel_pc,#js_tags_preview_toast{
  display:none!important;
}
</style>"""


def _visible_html_text(html: str) -> str:
    cleaned = _SCRIPT_RE.sub(" ", html or "")
    cleaned = _STYLE_BLOCK_RE.sub(" ", cleaned)
    text = unescape(_HTML_TAG_RE.sub(" ", cleaned))
    return re.sub(r"\s+", " ", text).strip()


def html_has_static_preview(html: str) -> bool:
    source = html or ""
    hint = _ARTICLE_HINT_RE.search(source)
    if hint:
        chunk = source[hint.start() : hint.start() + 200000]
        chunk_text = _visible_html_text(chunk)
        if re.search(r"<(?:p|img|h1|h2|h3|section)\b", chunk, re.I) and len(chunk_text) >= 4:
            return True
        if len(chunk_text) >= 80:
            return True
    vis = _visible_html_text(source)
    _, _, embedded = _embedded_article(source)
    if embedded:
        body = _html_to_text(embedded) if embedded.lstrip().startswith("<") else embedded
        sample = re.sub(r"\s+", "", body)[:48]
        if len(sample) >= 12 and sample not in re.sub(r"\s+", "", vis):
            return False
    if len(vis) < 120:
        return False
    if _SPA_SHELL_RE.search(source):
        return False
    return True


def _promote_lazy_images(html: str) -> str:
    def repl(match: re.Match[str]) -> str:
        tag = match.group(0)
        data = re.search(r"""\bdata-(?:src|original)=(["'])(.*?)\1""", tag, re.I)
        if not data:
            return tag
        url = unescape(data.group(2)).strip()
        if not url:
            return tag
        quoted = escape(url, quote=True)
        src = re.search(r"""(?<![\w-])src=(["'])(.*?)\1""", tag, re.I)
        if src:
            current = unescape(src.group(2)).strip()
            if current and not current.startswith("data:") and current != "about:blank":
                return tag
            return tag[: src.start()] + f'src="{quoted}"' + tag[src.end() :]
        return re.sub(r"^<img\b", f'<img src="{quoted}"', tag, count=1, flags=re.I)

    return _IMG_TAG_RE.sub(repl, html)


def _inject_head(html: str, snippet: str) -> str:
    close = re.search(r"</head>", html, re.I)
    if close:
        return html[: close.start()] + snippet + html[close.start() :]
    open_tag = re.search(r"<head[^>]*>", html, re.I)
    if open_tag:
        return html[: open_tag.end()] + snippet + html[open_tag.end() :]
    return snippet + html


def original_webpage_preview(path: Path, base_url: str = "") -> bytes | None:
    try:
        html = path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        return None
    if not html_has_static_preview(html):
        return None
    cleaned = _promote_lazy_images(_SCRIPT_RE.sub("", html))
    parts = [_ORIGINAL_PREVIEW_STYLE]
    href = (base_url or "").strip()
    if href.startswith("http://") or href.startswith("https://"):
        parts.insert(0, f'<base href="{escape(href, quote=True)}" />')
    return _inject_head(cleaned, "".join(parts)).encode("utf-8")


def embedded_article_preview_html(html: str, title: str = "") -> bytes | None:
    found_title, _, content = _embedded_article(html or "")
    body = (content or "").strip()
    if not body or not body.lstrip().startswith("<"):
        return None
    if len(_html_to_text(body)) < 12:
        return None
    body = _promote_lazy_images(_SCRIPT_RE.sub("", body))
    page_title = (found_title or title or "").strip() or "预览"
    return (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\" />"
        f"<title>{escape(page_title)}</title>"
        "<meta name=\"referrer\" content=\"no-referrer\" />"
        "<style>body{max-width:42rem;margin:0 auto;padding:16px;"
        "font:16px/1.7 system-ui,'Songti SC',SimSun,sans-serif;"
        "color:#1f2937;background:#fff}"
        "img{max-width:100%;height:auto}h1{font-size:1.4rem;margin:0 0 1rem}</style>"
        f"</head><body><h1>{escape(page_title)}</h1>{body}</body></html>"
    ).encode("utf-8")


def segments_preview_html(title: str, segments: list[TranscriptSegment]) -> bytes:
    blocks = []
    for item in segments:
        text = escape(item.text).replace("\n", "<br />")
        blocks.append(f'<p id="seg-{item.id}">{text}</p>')
    body = "\n".join(blocks) or "<p>没有可预览的文字</p>"
    return (
        "<!doctype html><html lang=\"zh-CN\"><head><meta charset=\"utf-8\" />"
        f"<title>{escape(title)}</title>"
        "<style>body{font:16px/1.65 system-ui,sans-serif;margin:0;padding:16px;"
        "color:#1f2937;background:#fff}p{margin:0 0 12px}</style></head><body>"
        f"{body}</body></html>"
    ).encode("utf-8")


def webpage_preview_html(
    path: Path,
    title: str = "",
    transcript_json: str = "",
    source_url: str = "",
) -> bytes:
    original = original_webpage_preview(path, base_url=source_url)
    if original:
        return original
    try:
        raw_html = path.read_text(encoding="utf-8", errors="ignore")
    except OSError:
        raw_html = ""
    embedded = embedded_article_preview_html(raw_html, title=title)
    if embedded:
        return embedded
    segments: list[TranscriptSegment] = []
    if (transcript_json or "").strip():
        from app.services.jsonutil import loads

        raw = loads(transcript_json, [])
        if isinstance(raw, list):
            segments = [TranscriptSegment.model_validate(item) for item in raw]
    if not segments:
        extracted = extract_html_bytes(path.read_bytes(), fallback_title=title or path.stem)
        segments = extracted.segments
        title = title or extracted.title
    if not segments:
        raise DocumentError("未能从网页提取到正文")
    return segments_preview_html(title or path.stem, segments)


def office_preview_html(path: Path) -> bytes:
    cache = path.parent / "preview.html"
    try:
        if cache.is_file() and cache.stat().st_mtime >= path.stat().st_mtime and cache.stat().st_size > 0:
            return cache.read_bytes()
    except OSError:
        pass
    docx_path = path
    suffix = path.suffix.lower()
    if suffix == ".doc":
        sibling = path.with_suffix(".docx")
        if sibling.is_file() and sibling.stat().st_size > 0:
            docx_path = sibling
        else:
            from app.services.plugins import convert_doc_with_soffice

            docx_path = convert_doc_with_soffice(path, path.parent)
    extracted = extract_docx(docx_path)
    html = segments_preview_html(extracted.title or path.stem, extracted.segments)
    cache.write_bytes(html)
    return html
